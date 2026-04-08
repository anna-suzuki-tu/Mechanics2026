"""
Supplementary Video Framework v2
- 6-slide structure matching the original 1-3 design
- Slide 1: Hook (scene setup with visual placeholder)
- Slide 2: Contradiction (question pose)
- Slide 3: Correction statement (with key visual if provided)
- Slide 4: Definition / Formula
- Slide 5: Contrast (wrong = right, two-column)
- Slide 6: Summary takeaway
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from lxml import etree
import os

# ── Canvas ────────────────────────────────────────────────
VIDEO_W = Inches(4.21875)   # 9:16
VIDEO_H = Inches(7.5)

# ── Margins / Zones ───────────────────────────────────────
MX = Inches(0.3)            # horizontal margin
INNER_W = VIDEO_W - 2 * MX  # usable width

TOP_ZONE_Y  = Inches(0.35)
TOP_ZONE_H  = Inches(1.0)

MID_ZONE_Y  = Inches(1.45)
MID_ZONE_H  = Inches(3.8)

BOT_ZONE_Y  = Inches(5.35)
BOT_ZONE_H  = Inches(1.8)

LABEL_Y = VIDEO_H - Inches(0.42)
LABEL_H = Inches(0.35)

# ── Palette (matches original 1-3) ────────────────────────
BG      = RGBColor(0x46, 0x78, 0x9B)   # steel blue background
WHITE   = RGBColor(0xF0, 0xF0, 0xF0)   # near-white text
GOLD    = RGBColor(0xC6, 0xA5, 0x4A)   # gold accent
DIM     = RGBColor(0xC8, 0xC8, 0xC8)   # muted label
NEQ_RED = RGBColor(0xFF, 0x6B, 0x6B)   # red for wrong side

FONT = "Helvetica Neue"


# ─────────────────────────────────────────────────────────
class SupplementaryVideoFramework:

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width  = VIDEO_W
        self.prs.slide_height = VIDEO_H

    # ── Low-level helpers ─────────────────────────────────

    def _bg(self, slide):
        sh = slide.shapes.add_shape(
            1,  # MSO_SHAPE.RECTANGLE
            0, 0, VIDEO_W, VIDEO_H
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = BG
        sh.line.fill.background()
        # Push behind everything
        slide.shapes._spTree.remove(sh._element)
        slide.shapes._spTree.insert(2, sh._element)

    def _tb(self, slide, text, x, y, w, h,
            size=28, bold=False, color=WHITE,
            align=PP_ALIGN.CENTER, wrap=True):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf  = box.text_frame
        tf.word_wrap = wrap
        tf.clear()
        p   = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.name  = FONT
        run.font.size  = Pt(size)
        run.font.bold  = bold
        run.font.color.rgb = color
        return box

    def _multiline_tb(self, slide, lines, x, y, w, h,
                      size=20, color=WHITE,
                      align=PP_ALIGN.LEFT):
        """lines: list of (text, bold, color) tuples or strings."""
        box = slide.shapes.add_textbox(x, y, w, h)
        tf  = box.text_frame
        tf.word_wrap = True
        tf.clear()
        for i, line in enumerate(lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.alignment = align
            if isinstance(line, tuple):
                txt, bld, clr = line
            else:
                txt, bld, clr = line, False, color
            run = p.add_run()
            run.text = txt
            run.font.name  = FONT
            run.font.size  = Pt(size)
            run.font.bold  = bld
            run.font.color.rgb = clr

    def _placeholder_box(self, slide, label="[visual]"):
        """Outlined placeholder rectangle in the mid zone."""
        sh = slide.shapes.add_shape(
            1, MX, MID_ZONE_Y, INNER_W, MID_ZONE_H
        )
        sh.fill.background()
        sh.line.color.rgb = GOLD
        sh.line.width = Pt(1.5)
        self._tb(slide, label,
                 MX, MID_ZONE_Y, INNER_W, MID_ZONE_H,
                 size=18, color=DIM)

    def _label(self, slide, text):
        self._tb(slide, text,
                 MX, LABEL_Y, INNER_W, LABEL_H,
                 size=12, color=DIM)

    def _new_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._bg(slide)
        return slide

    # ── Slide templates ───────────────────────────────────

    def slide_hook(self, title, label, graph_desc=""):
        """S1: Hook — title top, visual placeholder below with clear gap."""
        s = self._new_slide()
        # Title box tall enough for 3 lines at 26pt
        self._tb(s, title,
                 MX, Inches(0.3), INNER_W, Inches(2.0),
                 size=26, bold=True, color=WHITE)
        # Placeholder box: starts after title zone with clear gap
        ph_y = Inches(2.5)
        ph_h = VIDEO_H - ph_y - Inches(0.55)
        sh = s.shapes.add_shape(1, MX, ph_y, INNER_W, ph_h)
        sh.fill.background()
        sh.line.color.rgb = GOLD
        sh.line.width = Pt(1.5)
        if graph_desc:
            self._tb(s, graph_desc,
                     MX + Inches(0.1), ph_y + Inches(0.15),
                     INNER_W - Inches(0.2), ph_h - Inches(0.3),
                     size=14, color=DIM)
        self._label(s, label)

    def slide_question(self, question, label, sub=""):
        """S2: Pose the misconception as a question.
        Auto-size font based on question length to avoid overflow."""
        s = self._new_slide()
        # Adaptive font size: shorter text = larger font
        q_len = len(question)
        if q_len < 55:
            q_size, q_h = 32, Inches(2.6)
        elif q_len < 85:
            q_size, q_h = 26, Inches(3.0)
        else:
            q_size, q_h = 20, Inches(3.4)
        self._tb(s, question,
                 MX, Inches(0.7), INNER_W, q_h,
                 size=q_size, bold=True, color=WHITE)
        if sub:
            sub_y = Inches(0.7) + q_h + Inches(0.25)
            sub_h = min(Inches(1.3), VIDEO_H - sub_y - LABEL_H - Inches(0.3))
            self._tb(s, sub,
                     MX, sub_y, INNER_W, sub_h,
                     size=14, bold=False, color=GOLD)
        self._label(s, label)

    def slide_correction(self, statement, formula, label):
        """S3: Reveal the correct concept + formula. No overlap."""
        s = self._new_slide()
        # Statement: capped at 1.6" — at most 2 lines at 22pt
        self._tb(s, statement,
                 MX, Inches(0.4), INNER_W, Inches(1.6),
                 size=22, bold=True, color=GOLD)
        # Formula: starts at 2.2" — clear gap below statement
        self._tb(s, formula,
                 MX, Inches(2.2), INNER_W, Inches(3.2),
                 size=32, bold=True, color=WHITE)
        self._label(s, label)

    def slide_definition(self, term, definition, label):
        """S4: Big formula / definition."""
        s = self._new_slide()
        self._tb(s, term,
                 MX, Inches(1.2), INNER_W, Inches(1.0),
                 size=22, bold=False, color=DIM)
        self._tb(s, "||",
                 MX, Inches(2.3), INNER_W, Inches(0.6),
                 size=26, bold=True, color=GOLD)
        self._tb(s, definition,
                 MX, Inches(3.0), INNER_W, Inches(1.4),
                 size=26, bold=False, color=WHITE)
        self._label(s, label)

    def slide_contrast(self, title,
                       wrong_top, wrong_bot,
                       right_top, right_bot,
                       label):
        """S5: Wrong vs. Correct two-column contrast."""
        s  = self._new_slide()
        gap = Inches(0.12)
        cw  = (INNER_W - gap) / 2
        lx  = MX
        rx  = MX + cw + gap
        cy  = Inches(1.55)

        self._tb(s, title,
                 MX, Inches(0.28), INNER_W, Inches(1.1),
                 size=24, bold=True, color=WHITE)

        # Divider line between columns
        line = s.shapes.add_shape(1,
            MX + cw + gap/2 - Pt(0.5),
            cy, Pt(1), Inches(4.5))
        line.fill.solid()
        line.fill.fore_color.rgb = DIM
        line.line.fill.background()

        # Left (wrong) ─ muted white
        self._tb(s, wrong_top,
                 lx, cy, cw, Inches(0.6),
                 size=16, bold=True, color=DIM,
                 align=PP_ALIGN.CENTER)
        self._tb(s, wrong_bot,
                 lx, cy + Inches(0.65), cw, Inches(3.5),
                 size=15, bold=False, color=WHITE,
                 align=PP_ALIGN.LEFT)

        # Right (correct) ─ gold accent
        self._tb(s, right_top,
                 rx, cy, cw, Inches(0.6),
                 size=16, bold=True, color=GOLD,
                 align=PP_ALIGN.CENTER)
        self._tb(s, right_bot,
                 rx, cy + Inches(0.65), cw, Inches(3.5),
                 size=15, bold=False, color=GOLD,
                 align=PP_ALIGN.LEFT)

        self._label(s, label)

    def slide_summary(self, takeaway, label):
        """S6: One-line takeaway."""
        s = self._new_slide()
        self._tb(s, takeaway,
                 MX, Inches(2.5), INNER_W, Inches(2.0),
                 size=30, bold=True, color=WHITE)
        self._label(s, label)

    # ── High-level builder ────────────────────────────────

    def build_video(self, v, hook_img=None, contradiction_img=None):
        lbl = f"Supplementary Video {v['id']}"

        # S1 Hook
        self.slide_hook(
            title      = v['hook'],
            label      = lbl,
            graph_desc = v.get('hook_graph', '')
        )

        # S2 Contradiction / Question
        self.slide_question(
            question = v['contradiction'],
            label    = lbl,
            sub      = v.get('contradiction_graph', '')
        )

        # S3 Correction
        self.slide_correction(
            statement = v['correction_statement'],
            formula   = v.get('correction_formula', ''),
            label     = lbl
        )

        # S4 Definition (built from correction_formula + contrast info)
        # Use contrast title as the 'term' and script[3] as the 'definition'
        script = v.get('script', ('',)*5)
        self.slide_definition(
            term       = v['contrast_title'],
            definition = script[3] if len(script) > 3 else '',
            label      = lbl
        )

        # S5 Contrast
        self.slide_contrast(
            title      = v['contrast_title'],
            wrong_top  = v['contrast_left_label'],
            wrong_bot  = v['contrast_left_body'],
            right_top  = v['contrast_right_label'],
            right_bot  = v['contrast_right_body'],
            label      = lbl
        )

        # S6 Summary
        self.slide_summary(
            takeaway = v['summary'],
            label    = lbl
        )

    def save(self, path):
        os.makedirs(os.path.dirname(path) if os.path.dirname(path) else '.', exist_ok=True)
        self.prs.save(path)


# ── CLI test ─────────────────────────────────────────────
if __name__ == "__main__":
    import sys
    sys.path.insert(0, '.')
    from video_data import VIDEOS

    vid_id = sys.argv[1] if len(sys.argv) > 1 else "1-3"
    v = next(x for x in VIDEOS if x['id'] == vid_id)
    fw = SupplementaryVideoFramework()
    fw.build_video(v)
    out = f"test_{vid_id.replace('-','_')}.pptx"
    fw.save(out)
    print(f"Saved: {out}")
